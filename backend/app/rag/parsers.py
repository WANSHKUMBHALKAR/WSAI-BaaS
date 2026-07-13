"""
Document parsers for PDF, DOCX, plain text, and web URLs.
All parsers return a plain string of extracted text.
"""
from __future__ import annotations

import io
import logging
from pathlib import Path
from typing import Optional

import httpx

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def parse_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF using PyMuPDF (fitz)."""
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        pages = [page.get_text() for page in doc]
        return "\n".join(pages)
    except ImportError:
        # Fallback: pypdf
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def parse_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX file using python-docx."""
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also extract tables
    for table in doc.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(paragraphs)


# ---------------------------------------------------------------------------
# Plain Text
# ---------------------------------------------------------------------------
def parse_text(file_bytes: bytes, encoding: str = "utf-8") -> str:
    return file_bytes.decode(encoding, errors="replace")


# ---------------------------------------------------------------------------
# URL / Website
# ---------------------------------------------------------------------------
async def parse_url(url: str) -> str:
    """Download and strip HTML from a webpage."""
    try:
        from bs4 import BeautifulSoup
        async with httpx.AsyncClient(follow_redirects=True, timeout=30.0) as client:
            resp = await client.get(url)
            resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "html.parser")
        # Remove scripts, styles, nav, footer
        for tag in soup(["script", "style", "nav", "footer", "header"]):
            tag.decompose()
        return soup.get_text(separator="\n", strip=True)
    except Exception as exc:
        logger.error("URL parsing failed for %s: %s", url, exc)
        raise


# ---------------------------------------------------------------------------
# Dispatcher
# ---------------------------------------------------------------------------
async def parse_document(
    file_bytes: Optional[bytes],
    filename: str,
    url: Optional[str] = None,
) -> str:
    if url:
        return await parse_url(url)

    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(file_bytes)
    elif suffix in (".docx", ".doc"):
        return parse_docx(file_bytes)
    else:
        return parse_text(file_bytes)
